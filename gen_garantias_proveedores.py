# -*- coding: utf-8 -*-
"""Documento independiente: garantías de los proveedores de modelos (fuentes oficiales)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

FONT = 'Arial'
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x05, 0x63, 0xC1)
HEADER_FILL = "D9D9D9"
ZEBRA = "F2F2F2"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

normal = doc.styles['Normal']
normal.font.name = FONT; normal.font.size = Pt(11); normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def add_hyperlink(paragraph, url, text, size=9, color=LINK):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:ascii'), FONT); rf.set(qn('w:hAnsi'), FONT); rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size*2))); rPr.append(sz)
    col = OxmlElement('w:color'); col.set(qn('w:val'), '%02X%02X%02X' % (color[0], color[1], color[2])); rPr.append(col)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)
    return hyperlink

def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = BLACK; r.font.name = FONT
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),'000000')
    pbdr.append(bottom); pPr.append(pbdr)

def para(text, size=11, italic=False, bold=False, after=6, color=BLACK):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic; r.bold = bold
    r.font.name = FONT; r.font.color.rgb = color
    return p

def src_table(rows):
    """rows: list of (garantia, label, url)."""
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, htext in enumerate(['Garantía', 'Fuente oficial']):
        hdr[i].text = ''; pp = hdr[i].paragraphs[0]
        rr = pp.add_run(htext); rr.bold = True; rr.font.size = Pt(10); rr.font.name = FONT; rr.font.color.rgb = BLACK
        shade(hdr[i], HEADER_FILL)
    for idx, (gar, label, url) in enumerate(rows):
        cells = t.add_row().cells
        cells[0].text = ''; p0 = cells[0].paragraphs[0]
        r0 = p0.add_run(gar); r0.font.size = Pt(10); r0.font.name = FONT; r0.font.color.rgb = BLACK
        cells[1].text = ''; p1 = cells[1].paragraphs[0]
        rl = p1.add_run(label + ' — '); rl.font.size = Pt(9.5); rl.font.name = FONT; rl.font.color.rgb = BLACK
        add_hyperlink(p1, url, url, size=8.5)
        if idx % 2 == 0:
            shade(cells[0], ZEBRA); shade(cells[1], ZEBRA)
    for row in t.rows:
        row.cells[0].width = Cm(6.3); row.cells[1].width = Cm(10.7)
    return t

# ===================== PORTADA / CABECERA =====================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run('Asistente Inteligente de Apoyo al Médico — Guías VIH')
r.font.size = Pt(10.5); r.font.color.rgb = GREY; r.font.name = FONT
ph = doc.add_paragraph()
r = ph.add_run('Garantías de los proveedores de modelos de IA')
r.bold = True; r.font.size = Pt(19); r.font.color.rgb = BLACK; r.font.name = FONT
ps = doc.add_paragraph(); ps.paragraph_format.space_after = Pt(10)
r = ps.add_run('Fuentes oficiales para la verificación en materia de protección de datos (RGPD)')
r.italic = True; r.font.size = Pt(11.5); r.font.color.rgb = GREY; r.font.name = FONT

para('Documento de referencia preparado para la asesoría jurídica y el Delegado de Protección de Datos '
     '(DPD) de la organización. Recopila los enlaces oficiales donde se pueden verificar las garantías '
     'de privacidad y seguridad de los dos proveedores de modelos contemplados en la propuesta (Amazon '
     'Bedrock y Azure OpenAI). No constituye asesoramiento legal. Complementa el Anexo A de la propuesta '
     '(apartado A.8, «Garantías de los proveedores de modelos»). Enlaces verificados a fecha de junio de '
     '2026; conviene confirmar la versión vigente de cada documento contractual.', size=10.5, color=GREY)

# ===================== AMAZON BEDROCK =====================
h1('1. Amazon Bedrock (modelos Claude)')
src_table([
    ('No se entrena con los datos del cliente · no se comparten con los proveedores de modelos · no se retienen los prompts',
     'Cómo usa Amazon Bedrock los datos de entrada y salida (AWS re:Post / Knowledge Center)',
     'https://repost.aws/knowledge-center/amazon-bedrock-model-data-use'),
    ('Resumen oficial de seguridad, privacidad e IA responsable',
     'Amazon Bedrock — Security, Privacy & Responsible AI',
     'https://aws.amazon.com/bedrock/security-privacy-responsible-ai/'),
    ('Detalle técnico: cuentas de despliegue aisladas por proveedor, cifrado en tránsito y reposo, KMS, PrivateLink',
     'Amazon Bedrock — Security overview (documentación)',
     'https://docs.aws.amazon.com/bedrock/latest/userguide/security-overview.html'),
    ('Preguntas frecuentes (sección de seguridad y privacidad)',
     'Amazon Bedrock — FAQs',
     'https://aws.amazon.com/bedrock/faqs/'),
    ('Acuerdo de tratamiento de datos (DPA) conforme al RGPD, con cláusulas contractuales tipo',
     'AWS GDPR Data Processing Addendum (PDF)',
     'https://d1.awsstatic.com/legal/aws-gdpr/aws-gdpr-dpa-online.pdf'),
    ('Centro de cumplimiento RGPD',
     'AWS GDPR Center',
     'https://aws.amazon.com/compliance/gdpr-center/'),
    ('Residencia de datos en la UE y soberanía digital europea',
     'AWS EU Data Protection · European Digital Sovereignty FAQ',
     'https://aws.amazon.com/compliance/eu-data-protection/'),
])

# ===================== AZURE OPENAI =====================
h1('2. Azure OpenAI (Microsoft)')
src_table([
    ('No se entrena con tus datos · no se comparten con OpenAI · modelos sin estado (no se almacenan prompts)',
     'Data, privacy, and security for Azure OpenAI / Models sold by Azure',
     'https://learn.microsoft.com/en-us/azure/ai-foundry/responsible-ai/openai/data-privacy'),
    ('Monitorización de abuso y exclusión voluntaria («modified abuse monitoring» opt-out) y cómo verificar que está desactivada',
     'Azure OpenAI — Abuse monitoring',
     'https://learn.microsoft.com/en-us/azure/ai-foundry/openai/concepts/abuse-monitoring'),
    ('Acuerdo de tratamiento de datos (DPA) conforme al RGPD, con cláusulas contractuales tipo',
     'Microsoft Products and Services Data Protection Addendum (DPA)',
     'https://aka.ms/DPA'),
    ('Términos de privacidad y seguridad (Product Terms)',
     'Microsoft — Privacy and Security Terms',
     'https://www.microsoft.com/licensing/terms/product/PrivacyandSecurityTerms/all'),
    ('Residencia de datos en la UE (EU Data Boundary; despliegues DataZone UE)',
     'Microsoft — EU Data Boundary',
     'https://learn.microsoft.com/en-us/privacy/eudb/eu-data-boundary-learn'),
    ('Centro de confianza (compromisos de privacidad y seguridad)',
     'Microsoft Trust Center',
     'https://www.microsoft.com/trust-center'),
])

# ===================== NOTAS =====================
h1('3. Notas para la revisión jurídica')
for t in [
 'La página de datos y privacidad de Azure OpenAI confirma que los prompts, respuestas, embeddings y datos '
 'de entrenamiento no se ponen a disposición de OpenAI ni se usan para entrenar modelos sin permiso, y que '
 'los modelos son «sin estado» (no almacenan prompts).',
 'La exclusión de la monitorización de abuso («modified abuse monitoring»), recomendable cuando se traten '
 'datos sensibles, requiere aprobación de Microsoft y se solicita mediante el formulario enlazado desde la '
 'documentación de «Abuse monitoring».',
 'En Amazon Bedrock, el DPA global incorpora las cláusulas contractuales tipo y aplica automáticamente '
 'cuando el cliente está sujeto al RGPD; debe confirmarse la región de la UE concreta elegida para el '
 'despliegue.',
 'Para una residencia estricta en la UE: en Bedrock, fijar la región de la UE; en Azure OpenAI, emplear '
 'despliegues «DataZone UE» (no «Global»), ya que los despliegues «Global» pueden procesar en cualquier '
 'geografía.',
 'La elección de proveedor (Bedrock o Azure OpenAI) no altera el modelo de protección de datos de la '
 'solución: en ambos casos el modelo solo recibe datos ya seudonimizados y se invoca en la UE.',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(t); r.font.size = Pt(10); r.font.name = FONT; r.font.color.rgb = BLACK

doc.save('Garantias_Proveedores_Modelos_RGPD.docx')
print('OK -> Garantias_Proveedores_Modelos_RGPD.docx')
