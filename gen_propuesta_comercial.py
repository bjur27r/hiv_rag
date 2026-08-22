# -*- coding: utf-8 -*-
"""Genera la propuesta del Asistente VIH (tono sobrio, Arial negro, con anexo RGPD)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Arial'
BLACK = RGBColor(0x00, 0x00, 0x00)
HEADER_FILL = "D9D9D9"
ZEBRA = "F2F2F2"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)

normal = doc.styles['Normal']
normal.font.name = FONT; normal.font.size = Pt(11); normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.15

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def set_cell_text(cell, text, bold=False, size=10, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align: p.alignment = align
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.name = FONT; r.font.color.rgb = BLACK
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = BLACK; r.font.name = FONT
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6'); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),'000000')
    pbdr.append(bottom); pPr.append(pbdr)
    return p

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = BLACK; r.font.name = FONT
    return p

def para(text, size=11, italic=False, bold=False, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.italic = italic; r.bold = bold
    r.font.name = FONT; r.font.color.rgb = BLACK
    return p

def bullet(text, bold_lead=None, size=11):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=BLACK
    r2 = p.add_run(text); r2.font.name=FONT; r2.font.size=Pt(size); r2.font.color.rgb=BLACK
    return p

def numbered(text, size=11):
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=BLACK
    return p

def make_table(headers, rows, widths=None, zebra=True, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        set_cell_text(hdr[i], htext, bold=True, size=size); shade(hdr[i], HEADER_FILL)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            set_cell_text(cells[ci], val, size=size)
            if zebra and ri % 2 == 0: shade(cells[ci], ZEBRA)
    if widths:
        for row in t.rows:
            for ci, w in enumerate(widths):
                row.cells[ci].width = Cm(w)
    return t

def spacer(pts=6):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(pts)

# ========================= DIAGRAMA DE ARQUITECTURA =========================
def make_arch_diagram(path):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

    fig, ax = plt.subplots(figsize=(8.6, 10.2))
    ax.set_xlim(0, 100); ax.set_ylim(0, 122); ax.axis('off')
    EDGE = '#000000'

    def box(cx, cy, w, h, lines, fill='#FFFFFF', fs=10, bold_first=True, align='center'):
        ax.add_patch(FancyBboxPatch((cx-w/2, cy-h/2), w, h,
                     boxstyle='round,pad=0.6,rounding_size=2.2',
                     linewidth=1.3, edgecolor=EDGE, facecolor=fill))
        if isinstance(lines, str): lines = [lines]
        n = len(lines); lh = h/(n+0.6)
        y0 = cy + (n-1)*lh/2
        for i, ln in enumerate(lines):
            ax.text(cx, y0 - i*lh, ln, ha=align, va='center',
                    fontsize=(fs+1 if (i==0 and bold_first) else fs-0.5),
                    fontweight=('bold' if (i==0 and bold_first) else 'normal'),
                    color=EDGE, family='sans-serif')

    def varrow(y1, y2, x=50, two=True):
        ax.add_patch(FancyArrowPatch((x, y1), (x, y2), arrowstyle=('<|-|>' if two else '-|>'),
                     mutation_scale=15, linewidth=1.3, color=EDGE))

    # Capa transversal (derecha)
    ax.add_patch(FancyBboxPatch((84, 12), 14, 96, boxstyle='round,pad=0.6,rounding_size=2.2',
                 linewidth=1.2, edgecolor=EDGE, facecolor='#F2F2F2'))
    ax.text(91, 60, 'Seguridad y auditoría', ha='center', va='center', rotation=90,
            fontsize=10.5, fontweight='bold', color=EDGE)
    ax.text(91, 60, '\n\n\nCifrado · IAM · red privada (UE)   ·   registro auditable de cada consulta',
            ha='center', va='center', rotation=90, fontsize=8.2, color=EDGE)

    CW = 74; CX = 41  # ancho y centro de las cajas principales
    box(CX, 116, 30, 7, ['Médico especialista'], fill='#E8E8E8')
    box(CX, 104, CW, 9, ['Aplicación web — interfaz de consulta',
                         'Integrada en los dominios de la organización · historial · exportación citada'])
    box(CX, 91, CW, 9, ['API · Pasarela de seguridad',
                        'Autenticación · control de tráfico · seudonimización en la entrada · trazabilidad'])
    box(CX, 63, CW, 30, ['Motor agéntico (orquestador)',
                         'Interpretación y reescritura de la consulta',
                         'Recuperación híbrida (semántica + léxica)',
                         'Filtro de vigencia y reordenación',
                         'Generación de la respuesta con cita',
                         'Verificación de fidelidad (guardarraíl de salida)',
                         'Supervisión médica (confirmación ante datos sensibles)'], fs=9.5)
    box(CX, 38, CW, 10, ['Base de conocimiento',
                         'Guías estructuradas y fragmentadas · índice vectorial + léxico',
                         'Control de versiones y vigencia'])
    box(CX, 22, CW, 9, ['Capa de modelos de IA',
                        'Amazon Bedrock (Claude) / Azure OpenAI · embeddings · región UE'])

    varrow(112.5, 108.5); varrow(99.5, 95.5); varrow(86.5, 78)
    varrow(48, 43); varrow(33, 26.5)
    fig.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    fig.savefig(path, dpi=200, bbox_inches='tight', facecolor='white')
    plt.close(fig)

ARCH_IMG = 'arch_vih.png'
try:
    make_arch_diagram(ARCH_IMG)
    HAVE_DIAGRAM = True
except Exception as e:
    print('Aviso: no se pudo generar el diagrama:', e)
    HAVE_DIAGRAM = False

# ========================= PORTADA =========================
for _ in range(4): spacer(10)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('PROPUESTA TÉCNICA Y ECONÓMICA'); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=BLACK; r.font.name=FONT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Asistente Inteligente de Apoyo al Médico'); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=BLACK; r.font.name=FONT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('sobre las guías clínicas de VIH'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=BLACK; r.font.name=FONT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10)
r = p.add_run('Consulta basada en evidencia, con respuestas citadas y verificadas')
r.italic=True; r.font.size=Pt(12); r.font.color.rgb=BLACK; r.font.name=FONT
for _ in range(7): spacer(10)
for txt,sz in [('Preparado por: [Tu empresa]',12),('Dirigido a: [Cliente]',12),
               ('Versión 4.0 · Junio 2026 · Confidencial',10.5)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.font.size=Pt(sz); r.font.name=FONT; r.font.color.rgb=BLACK
doc.add_page_break()

# ========================= 1. INTRODUCCIÓN Y CONTEXTO =========================
h1('1. Introducción y contexto')
para('La atención a las personas que viven con VIH se apoya en guías de práctica clínica extensas y en '
     'constante actualización. Estas guías recogen, para cada situación, la recomendación más adecuada '
     'junto con su nivel de evidencia, y constituyen la referencia que el profesional debe seguir. Sin '
     'embargo, su volumen y la frecuencia con la que se revisan hacen que localizar la recomendación '
     'exacta, vigente y aplicable a un paciente concreto consuma un tiempo valioso en el momento de la '
     'consulta, con el riesgo añadido de trabajar con información ya desactualizada.')
para('Esta propuesta describe el desarrollo de un asistente para dar soporte al especialista en la '
     'consulta de las guías de forma natural: el profesional formula su pregunta en lenguaje corriente y '
     'el sistema responde con la recomendación pertinente, indicando la sección de la guía y el nivel de '
     'evidencia en que se basa. El objetivo no es sustituir el criterio clínico, sino ponerlo al '
     'servicio del profesional con mayor rapidez y trazabilidad, reduciendo la carga de búsqueda manual '
     'y el riesgo de trabajar con información no vigente.')
para('El asistente está pensado para integrarse en el entorno digital de la organización: se ofrece como '
     'un servicio accesible a través de una interfaz (API) y, sobre ella, una aplicación web alojada en '
     'los dominios de la propia organización. De este modo, la solución encaja con los sistemas y las '
     'políticas de seguridad ya existentes.')

# ========================= 2. DESCRIPCIÓN DEL SERVICIO =========================
h1('2. Descripción del servicio')
para('El servicio es un asistente conversacional especializado en las guías clínicas de VIH. El médico '
     'le plantea una duda —sobre una pauta de inicio, un ajuste por insuficiencia renal, una interacción '
     'farmacológica, el seguimiento de un paciente estable, etc.— y el asistente devuelve una respuesta '
     'redactada en lenguaje claro, acompañada de la referencia exacta a la guía (sección y página) y de '
     'su grado de recomendación.')
para('Para el médico hay una única forma de usarlo: preguntar. La diferencia está solo en cuánta '
     'información decide aportar. Puede formular la consulta en abstracto —qué recomienda la guía para '
     'iniciar tratamiento, qué precauciones existen ante una determinada interacción— o puede '
     'enriquecerla con la situación concreta de un paciente: sus condiciones, su medicación, sus '
     'analíticas. En este segundo caso el asistente ajusta la respuesta a ese perfil, indicando qué '
     'recomienda la guía y, cuando procede, qué opciones quedan descartadas y por qué. Cuanto más '
     'contexto recibe, más afinada es la respuesta, pero el modo de preguntar y de recibir la respuesta '
     'es siempre el mismo.')
para('Cuando la consulta incorpora datos de un paciente y el asistente detecta una contraindicación '
     'relevante o un dato dudoso, no continúa sin más: lo advierte y pide la confirmación del médico '
     'antes de seguir. De este modo, la herramienta nunca toma decisiones por su cuenta en las '
     'situaciones delicadas. En todos los casos, el asistente nunca presenta una recomendación sin '
     'indicar de dónde sale, para que el profesional pueda verificarla en segundos.')

# ========================= 3. EXPERIENCIA DE USUARIO OBJETIVO =========================
h1('3. Experiencia de usuario objetivo')
para('El destinatario es el médico especialista. La experiencia se diseña para que la herramienta '
     'resulte natural de usar en el día a día asistencial y para que genere confianza desde la primera '
     'consulta. El acceso se realiza a través de una aplicación web sencilla, integrada en los dominios '
     'de la organización, con historial de consultas y la posibilidad de exportar la respuesta citada '
     'cuando proceda incorporarla al curso clínico. Tres principios guían el diseño.')

h2('3.1. Trazabilidad inmediata')
para('Cada respuesta enlaza con la sección y la página concretas de la guía y muestra el grado de '
     'recomendación. El médico no recibe una opinión del sistema, sino un puntero verificable a la '
     'fuente oficial. Esto convierte la herramienta en un atajo hacia la guía, no en un sustituto opaco '
     'de ella.')

h2('3.2. Honestidad y abstención')
para('Cuando una pregunta no está cubierta por las guías, el asistente lo dice con claridad en lugar de '
     'improvisar una respuesta. Esta abstención explícita es una característica deliberada: en un '
     'contexto clínico, una respuesta inventada es mucho más peligrosa que un «esto no lo recoge la '
     'guía».')

h2('3.3. El médico siempre decide')
para('La herramienta asiste, nunca prescribe. En las consultas apoyadas en la situación del paciente, '
     'ante una contraindicación seria o un dato ambiguo, el flujo se detiene y solicita la confirmación '
     'del profesional antes de seguir. La decisión clínica final es, en todo momento, del médico.')

h2('3.4. Ejemplos de preguntas y respuestas')
para('Los siguientes ejemplos ilustran el formato de interacción y el estilo de respuesta. El contenido '
     'clínico concreto procede en cada caso de la guía vigente cargada en el sistema; aquí se muestran '
     'de forma esquemática para reflejar la estructura (respuesta + cita + nivel de evidencia) y los '
     'comportamientos de confirmación y abstención.', italic=True, size=10.5)
make_table(['Consulta del médico', 'Respuesta del asistente'],
    [['«¿Qué pauta de inicio recomienda la guía para un adulto sin tratamiento previo?»',
      'Resumen de la pauta preferente de inicio según la guía, con una breve justificación. '
      'Referencia: GeSIDA, sección «Tratamiento de inicio», pág. NN · Nivel de evidencia A-I.'],
     ['«Paciente en tratamiento estable con TDF/FTC/EFV que presenta ahora un filtrado glomerular de 45 mL/min. ¿Conviene ajustar?»',
      'El asistente advierte de la precaución asociada al tenofovir-DF en insuficiencia renal, muestra la '
      'recomendación de la guía sobre el ajuste o el cambio de pauta y, al tratarse de una situación '
      'delicada, solicita confirmación del médico antes de continuar. Referencia a la sección y nivel de '
      'evidencia correspondientes.'],
     ['«¿Qué dosis de [un fármaco no recogido en las guías de VIH] debo usar?»',
      '«Esta cuestión no está recogida en las guías de VIH cargadas; no dispongo de una recomendación '
      'citable al respecto.» (Abstención explícita en lugar de improvisar.)']],
    widths=[6.0, 11.0], size=10)
para('La aplicación web ofrece, además, interfaz de consulta clara con las citas enlazadas, acceso '
     'seguro integrado en los sistemas de identidad de la organización, historial de consultas, '
     'exportación de la respuesta citada al curso clínico y diseño adaptable a ordenador y a '
     'dispositivos móviles.')

# ========================= 4. ARQUITECTURA TENTATIVA =========================
h1('4. Arquitectura tentativa propuesta')
para('La solución se construye sobre un patrón ampliamente adoptado en sistemas de inteligencia '
     'artificial fiables (recuperación aumentada, RAG): en lugar de que el modelo responda «de memoria», '
     'primero se recuperan los fragmentos relevantes de las guías y, solo sobre ese material, se genera '
     'la respuesta. A ello se añade una verificación posterior que garantiza que todo lo afirmado está '
     'respaldado por la fuente, y un conjunto de guardarraíles que protegen tanto la calidad clínica '
     'como los datos personales.')

h2('4.1. Recorrido de una consulta')
para('De forma simplificada, cada consulta atraviesa las siguientes etapas:')
numbered('La pregunta pasa un primer filtro de seguridad y, si contiene datos del paciente, un proceso de seudonimización antes de cualquier otro tratamiento.')
numbered('El asistente interpreta la pregunta y, si conviene, la reformula para localizar mejor la información en las guías.')
numbered('Busca en las guías combinando dos técnicas: una que entiende el significado de la pregunta y otra que localiza términos exactos (fármacos, dosis, códigos).')
numbered('Selecciona los fragmentos más pertinentes y descarta cualquier material que no corresponda a la versión vigente de la guía.')
numbered('Redacta la respuesta apoyándose únicamente en esos fragmentos, citando la fuente de cada afirmación.')
numbered('Verifica que todo lo dicho está respaldado por la guía; si algo no lo está, lo descarta o avisa de que no se ha encontrado.')
numbered('Registra la consulta seudonimizada, la versión de guía utilizada y la respuesta, de modo que todo quede auditable.')

h2('4.2. Componentes arquitectónicos')
para('La solución se organiza en capas, cada una con una responsabilidad clara. El siguiente diagrama '
     'resume la arquitectura y el recorrido de una consulta, desde la aplicación web hasta la base de '
     'conocimiento y los modelos, con la seguridad y la auditoría como capa transversal.')
if HAVE_DIAGRAM:
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _AL
    pic_p = doc.add_paragraph(); pic_p.alignment = _AL.CENTER
    pic_p.paragraph_format.space_before = Pt(4); pic_p.paragraph_format.space_after = Pt(4)
    pic_p.add_run().add_picture(ARCH_IMG, width=Cm(14.0))
    cap = doc.add_paragraph(); cap.alignment = _AL.CENTER; cap.paragraph_format.space_after = Pt(8)
    rc = cap.add_run('Figura 1. Arquitectura por capas del asistente.')
    rc.italic = True; rc.font.size = Pt(9); rc.font.name = FONT; rc.font.color.rgb = BLACK
para('Los bloques principales son:')
bullet('punto de acceso del médico, alojada e integrada en los '
       'dominios de la organización y sujeta a sus sistemas de identidad y seguridad. Ofrece la '
       'conversación de consulta con las citas enlazadas, el historial y la exportación de la respuesta '
       'citada al curso clínico. Es la única pieza que ve el profesional; toda su lógica se apoya en la '
       'API.', bold_lead='Aplicación web (interfaz de consulta): ')
bullet('puerta de entrada única y segura al asistente. Autentica las '
       'aplicaciones y los usuarios autorizados, controla el tráfico y los límites de uso, registra cada '
       'llamada y aísla la infraestructura interna. Aquí se activa, además, el guardarraíl de '
       'seudonimización antes de cualquier otro tratamiento. Al exponerse como API, la misma lógica '
       'puede consumirse en el futuro desde otros sistemas (p. ej., la historia clínica electrónica).',
       bold_lead='API y pasarela de seguridad: ')
bullet('el cerebro del sistema. Coordina el flujo de cada consulta y '
       'decide qué módulos intervienen y en qué orden: interpreta y reescribe la pregunta, recupera los '
       'fragmentos relevantes de las guías, descarta lo no vigente, redacta la respuesta citada, '
       'verifica que todo lo afirmado está respaldado y, ante una situación delicada, detiene el flujo '
       'para pedir confirmación al médico. Sus módulos internos se detallan en el apartado 4.3.',
       bold_lead='Motor agéntico (orquestador): ')
bullet('el repositorio de las guías, ya estructuradas y '
       'fragmentadas, sobre el que se realiza la búsqueda. Mantiene un doble índice —vectorial '
       '(semántico) y léxico (términos exactos)— y un control de versiones que marca la vigencia de cada '
       'fragmento, de modo que nunca se responda con material derogado.', bold_lead='Base de conocimiento: ')
bullet('los modelos de lenguaje que generan y verifican la respuesta '
       'y los modelos de «embeddings» que sustentan la búsqueda semántica. Se invocan en puntos de '
       'acceso ubicados en la UE (Amazon Bedrock o Azure OpenAI; véase el apartado 9.3).',
       bold_lead='Capa de modelos de IA: ')
bullet('presentes en todas las capas: cifrado en tránsito y en '
       'reposo, control de acceso por identidad, red privada y un registro auditable —en su versión '
       'seudonimizada— de cada consulta.', bold_lead='Seguridad y auditoría (transversal): ')

h2('4.3. Componentes funcionales del motor')
para('Dentro del motor agéntico, el asistente se compone de módulos especializados, cada uno con una '
     'responsabilidad bien definida. Esta separación permite mantener, evaluar y mejorar cada pieza de '
     'forma independiente, y es clave para la fiabilidad del conjunto.')
make_table(['Componente', 'Función en el sistema'],
    [['Pasarela de seguridad y filtro de entrada', 'Puerta de entrada de toda consulta. Verifica que la petición procede de una aplicación y un usuario autorizados, aplica un primer filtro de contenido y activa el guardarraíl de seudonimización.'],
     ['Orquestador', 'Coordina el flujo: decide qué módulos intervienen y en qué orden, y gestiona las pausas para pedir confirmación al médico. Garantiza que ninguna etapa se salta.'],
     ['Interpretación y reescritura', 'Traduce la pregunta del médico al lenguaje de las guías: expande abreviaturas y sinónimos clínicos y genera variantes para que una duda coloquial encuentre la sección correcta.'],
     ['Motor de recuperación híbrida', 'Combina búsqueda semántica (significado) y búsqueda léxica (términos exactos: fármacos, dosis, códigos), evitando tanto omitir un término crítico como quedarse en coincidencias literales.'],
     ['Filtro de vigencia y reordenación', 'Descarta los fragmentos que no pertenecen a la versión vigente de la guía y reordena el resto por relevancia. Impide que una recomendación derogada llegue a la respuesta.'],
     ['Generación con cita', 'Redacta la respuesta apoyándose solo en los fragmentos seleccionados, con la referencia (sección y página) y el nivel de evidencia. Tiene prohibido aportar conocimiento propio.'],
     ['Verificador de fidelidad', 'Antes de entregar la respuesta, comprueba que cada afirmación está respaldada por la fuente. Si detecta algo sin respaldo, lo elimina, pide nueva redacción o convierte la respuesta en un «no se ha encontrado en las guías».'],
     ['Supervisión médica', 'Ante una contraindicación seria o un dato ambiguo, detiene el flujo y solicita al médico que confirme o complete la información. El sistema asiste; la decisión es del profesional.'],
     ['Registro de auditoría', 'Guarda, por cada consulta, la pregunta (ya seudonimizada), la versión de guías empleada, los fragmentos citados y la respuesta. Permite reconstruir por qué el asistente respondió lo que respondió.']],
    widths=[5.0, 12.0], size=9.5)

h2('4.4. Guardarraíles: calidad clínica y protección de datos')
para('Los guardarraíles son controles automáticos que el sistema aplica sin excepción en cada consulta. '
     'Son la base de la seguridad de la herramienta:')
bullet('comprueba que cada afirmación de la respuesta está respaldada por las '
       'guías y bloquea cualquier contenido inventado. Funciona a la salida, antes de entregar la '
       'respuesta al médico.', bold_lead='Guardarraíl de fidelidad: ')
bullet('trata la fecha de vigencia de cada recomendación como un filtro '
       'obligatorio, de modo que el sistema nunca responde con material que no pertenezca a la versión '
       'actual de la guía. Una guía caducada puede dar lugar a una recomendación peligrosa; por eso esta '
       'vigilancia es una de las garantías de seguridad más importantes.', bold_lead='Guardarraíl de vigencia: ')
bullet('cuando una consulta incorpora información de un paciente, ésta '
       'atraviesa un guardarraíl de seudonimización situado en la entrada, antes de que el modelo reciba '
       'dato alguno. Detecta los identificadores directos (nombre, número de historia, documento, fechas '
       'concretas, teléfonos) y los reemplaza por marcadores neutros; el resto del sistema trabaja solo '
       'con los datos clínicos estrictamente necesarios, nunca con la identidad de la persona. Además '
       'vigila las combinaciones de datos que podrían reidentificar a alguien y, ante un riesgo elevado, '
       'generaliza esos datos o se abstiene de procesar la consulta.', bold_lead='Guardarraíl de seudonimización: ')

# ========================= 5. SEGURIDAD, PRIVACIDAD Y DESPLIEGUE =========================
h1('5. Seguridad, privacidad y despliegue')
para('La solución se despliega íntegramente sobre infraestructura gestionada y segura en la nube, en '
     'región de la Unión Europea para favorecer la residencia de datos, con cifrado en tránsito y en '
     'reposo, control de acceso y registro de auditoría de cada consulta. El acceso se publica a través '
     'de una pasarela de seguridad (API Gateway) que autentica las aplicaciones y los usuarios '
     'autorizados, controla el tráfico y registra todas las llamadas.')
para('El tratamiento de datos de salud se ajusta al RGPD y se valida con el Delegado de Protección de '
     'Datos de la organización. El análisis detallado de cumplimiento en materia de protección de datos '
     '—pensado para su revisión por la asesoría jurídica— se incluye en el Anexo A de esta propuesta.')

# ========================= 6. ENTREGABLES =========================
h1('6. Entregables')
para('Al término del proyecto, la organización recibe:')
bullet('el asistente clínico funcionando, accesible de forma segura mediante API.', bold_lead='Servicio del asistente: ')
bullet('desplegada e integrada en los dominios de la organización.', bold_lead='Aplicación web: ')
bullet('el conjunto de guías cargado, estructurado y con control de versiones.', bold_lead='Base de conocimiento: ')
bullet('documentación funcional, técnica, manual de usuario, documentación para '
       'integradores e informe de validación clínica.', bold_lead='Documentación completa: ')
bullet('la batería de preguntas y los resultados obtenidos, como evidencia de '
       'la calidad del sistema.', bold_lead='Batería de evaluación: ')

# ========================= 7. PLANIFICACIÓN =========================
h1('7. Planificación y entregables por fase')
para('El proyecto se organiza en siete fases sucesivas. Cada fase concluye con entregables concretos y '
     'verificables, lo que permite a la organización seguir el avance y validar resultados de forma '
     'progresiva. Duración total estimada: aproximadamente seis meses (26 semanas). El asistente queda '
     'utilizable por API en torno al tercer mes (MVP) y la aplicación web poco después.')
make_table(['Fase', 'Semanas', 'Actividades principales', 'Entregables'],
    [['1. Descubrimiento y diseño', '1–3',
      'Cierre de requisitos, criterios clínicos de aceptación, definición de la batería de evaluación y del modelo de datos.',
      'Documento de diseño funcional y plan de proyecto. Catálogo de criterios de aceptación.'],
     ['2. Base de conocimiento', '3–6',
      'Carga, estructuración, fragmentación y control de versiones de las guías; indexado híbrido.',
      'Base de conocimiento operativa y verificada, con trazabilidad de versiones.'],
     ['3. Motor del asistente (MVP)', '6–11',
      'Recuperación híbrida, reescritura de consulta, generación citada y verificación de fidelidad.',
      'Asistente funcional accesible por API (núcleo del MVP).'],
     ['4. Validación del MVP', '11–14',
      'Evaluación con la batería de preguntas clínicas reales y revisión por personal médico.',
      'Informe de validación y aceptación del MVP (hito de pago).'],
     ['5. Aplicación web y API segura', '14–19',
      'Pasarela API (autenticación, control de tráfico, trazabilidad) y aplicación web integrada en los dominios de la organización.',
      'Aplicación web desplegada y API documentada para integradores.'],
     ['6. Apoyo al paciente y privacidad', '19–24',
      'Consulta apoyada en la situación del paciente, supervisión médica y guardarraíl de seudonimización.',
      'Módulo validado y medidas de privacidad implantadas y documentadas.'],
     ['7. Despliegue y cierre', '24–26',
      'Puesta en producción, pruebas finales y entrega documental.',
      'Sistema en producción y documentación completa.']],
    widths=[3.6, 1.7, 6.4, 5.3], size=9.5)

# ========================= 8. INVERSIÓN =========================
h1('8. Inversión: desarrollo')
para('El desarrollo del asistente, su API y la aplicación web se presupuesta de forma cerrada. El '
     'detalle por paquete de trabajo se acompaña en el modelo económico adjunto; las cifras se ajustan '
     'al alcance final acordado en la fase de descubrimiento.')
make_table(['Bloque', 'Importe (sin IVA)'],
    [['Asistente clínico (núcleo) y su API', ''],
     ['Aplicación web y módulos complementarios', ''],
     ['TOTAL desarrollo', '']],
    widths=[11.5, 5.5])
h2('8.1. Plan de pagos')
para('Se propone un plan con una entrada reducida, un hito ligado a la validación del MVP y el resto '
     'repartido a lo largo de un año. El plan de pagos cubre únicamente el desarrollo.')
make_table(['Hito', 'Momento', '%', 'Importe'],
    [['Pago inicial', 'A la firma del contrato', '15 %', ''],
     ['Pago en MVP', 'Validación del MVP', '20 %', ''],
     ['Mensualidad (×12)', 'Meses 1–12 tras el MVP', '65 % /12', ''],
     ['Total desarrollo (sin IVA)', '', '100 %', '']],
    widths=[4.2, 6.0, 2.3, 4.5])

# ========================= 9. EXPLOTACIÓN =========================
h1('9. Costes de explotación')
para('Una vez en producción, el asistente requiere unos servicios en la nube para funcionar. Estos '
     'costes son recurrentes, dependen del uso real del sistema y se abonan directamente desde una '
     'cuenta de la propia organización; no forman parte del precio del desarrollo ni del mantenimiento. '
     'A continuación se detalla una estimación orientativa para un escenario de referencia.')

h2('9.1. Hipótesis de uso y consumo')
make_table(['Parámetro', 'Valor'],
    [['Médicos usuarios', '100'],
     ['Consultas por médico y día', '2'],
     ['Consultas diarias', '≈ 200'],
     ['Días laborables/mes', '≈ 22'],
     ['Consultas al mes', '≈ 4.400'],
     ['Tokens por consulta — generación citada (entrada / salida)', '≈ 12.000 / 800'],
     ['Tokens por consulta — reescritura + verificación (entrada / salida)', '≈ 8.000 / 500'],
     ['Volumen mensual — generación (entrada / salida)', '≈ 52,8 M / 3,5 M'],
     ['Volumen mensual — reescritura + verificación (entrada / salida)', '≈ 35,2 M / 2,2 M']],
    widths=[11.5, 5.5], size=9.5)
para('Ejemplo de consumo de una consulta: para la pregunta del apartado 3.4 sobre el paciente con '
     'filtrado glomerular de 45 mL/min, el sistema procesa la pregunta y los datos clínicos '
     'seudonimizados (~150 tokens), las instrucciones del asistente y las reglas de citación (~1.500 '
     'tokens) y los 6–8 fragmentos de guía recuperados (~10.000 tokens) —del orden de 12.000 tokens de '
     'entrada— y devuelve una respuesta citada de ~700–800 tokens; la verificación posterior relee la '
     'respuesta frente a esos fragmentos (~8.000 tokens de entrada). La estrategia de modelos emplea un '
     'modelo de calidad para redactar y un modelo más económico para verificar, y el almacenamiento en '
     'caché de prompts reduce el coste de la parte repetida de la entrada.', size=10, italic=True)

h2('9.2. Infraestructura base (configuración económica)')
para('La infraestructura se despliega en la nube, en región de la UE. Para minimizar el coste, la '
     'búsqueda se resuelve con un dominio gestionado de Amazon OpenSearch pequeño o con S3 Vectors, en '
     'lugar de una configuración de alta disponibilidad.')
make_table(['Componente', 'Servicio', 'Coste orientativo/mes'],
    [['Búsqueda (índice híbrido)', 'OpenSearch — dominio gestionado pequeño / S3 Vectors', '≈ 90–150 €'],
     ['Cómputo y orquestación', 'AWS Lambda / Fargate + Step Functions', '≈ 40–80 €'],
     ['Estado y sesiones', 'RDS pequeño / ElastiCache', '≈ 30–60 €'],
     ['Pasarela API', 'API Gateway (≈ 4.400 llamadas/mes)', '< 5 €'],
     ['Almacenamiento', 'S3 (guías y registros)', '≈ 5–10 €'],
     ['Observabilidad', 'CloudWatch + paneles', '≈ 20–40 €'],
     ['Seguridad y red', 'KMS, VPC, PrivateLink', '≈ 25–45 €'],
     ['Subtotal infraestructura', '', '≈ 215–390 €/mes']],
    widths=[4.6, 8.4, 4.0], size=9.5)
para('Nota: con Amazon OpenSearch Serverless en alta disponibilidad existe un mínimo de capacidad '
     'facturable (del orden de 2 OCU) que se paga aunque el sistema esté inactivo; en esa configuración '
     'la búsqueda sube a ≈ 330–700 €/mes. Para un corpus reducido como el de las guías, el dominio '
     'gestionado pequeño o S3 Vectors es la opción más económica.', size=9.5, italic=True)

h2('9.3. Capa de modelos: dos opciones')
para('La capa de modelos puede resolverse con Amazon Bedrock (modelos Claude) o con Azure OpenAI; en '
     'ambos casos los modelos se invocan en puntos de acceso ubicados en la UE. Costes mensuales '
     'orientativos para el volumen anterior (precios de referencia 2026; 1 USD ≈ 0,92 €).')
para('Opción A — Amazon Bedrock (modelos Claude). Para la generación se ofrecen dos niveles: una opción '
     'equilibrada (Sonnet) y una de máxima calidad (Opus). La verificación emplea siempre el modelo '
     'económico (Haiku). Precios por millón de tokens (entrada/salida): Opus 5/25 $, Sonnet 3/15 $, '
     'Haiku 1/5 $.', size=10)
make_table(['Tarea', 'Modelo', 'Tokens/mes (E / S)', 'Coste/mes'],
    [['Generación — opción equilibrada', 'Claude Sonnet 4.6', '52,8 M / 3,5 M', '≈ 194 €'],
     ['Generación — opción calidad', 'Claude Opus 4.8', '52,8 M / 3,5 M', '≈ 324 €'],
     ['Reescritura + verificación', 'Claude Haiku 4.5', '35,2 M / 2,2 M', '≈ 43 €'],
     ['Embeddings de consulta', 'Titan Embeddings', '≈ 0,9 M', '< 1 €'],
     ['Total modelos (Sonnet + Haiku)', '', '', '≈ 238 €/mes'],
     ['Total modelos (Opus + Haiku)', '', '', '≈ 368 €/mes']],
    widths=[4.8, 3.4, 4.4, 4.4], size=9.5)
para('Opción B — Azure OpenAI. Precios por millón de tokens (entrada/salida): GPT-4o 2,5/10 $, '
     'GPT-4o mini 0,15/0,60 $.', size=10)
make_table(['Tarea', 'Modelo', 'Tokens/mes (E / S)', 'Coste/mes'],
    [['Generación citada', 'GPT-4o', '52,8 M / 3,5 M', '≈ 154 €'],
     ['Reescritura + verificación', 'GPT-4o mini', '35,2 M / 2,2 M', '≈ 6 €'],
     ['Embeddings de consulta', 'text-embedding-3-small', '≈ 0,9 M', '< 1 €'],
     ['Total modelos — opción B', '', '', '≈ 161 €/mes']],
    widths=[4.8, 3.4, 4.4, 4.4], size=9.5)

h2('9.4. Total mensual orientativo')
make_table(['Escenario', 'Infraestructura', 'Modelos', 'Total orientativo/mes'],
    [['A — Bedrock (Sonnet) + OpenSearch económico', '≈ 215–390 €', '≈ 238 €', '≈ 450–630 €'],
     ['A — Bedrock (Opus) + OpenSearch económico', '≈ 215–390 €', '≈ 368 €', '≈ 580–760 €'],
     ['B — Azure OpenAI + OpenSearch económico', '≈ 215–390 €', '≈ 161 €', '≈ 375–550 €']],
    widths=[6.8, 3.4, 2.8, 4.0], size=9.5)
para('Cifras orientativas para ≈ 4.400 consultas/mes; deben confirmarse con la calculadora oficial del '
     'proveedor. Repartido entre las consultas, el coste de modelo equivale a unos 5 céntimos por '
     'consulta con la opción equilibrada y unos 8 céntimos con la opción de calidad; sumando la '
     'infraestructura, del orden de 10–17 céntimos por consulta. El coste de infraestructura es '
     'mayoritariamente fijo, de modo que el sistema escala bien: un aumento notable del número de '
     'médicos eleva el gasto solo de forma moderada.', size=9.5, italic=True)

# ========================= 10. MANTENIMIENTO Y SLA =========================
h1('10. Mantenimiento y niveles de servicio')
para('Se propone un contrato anual de mantenimiento que cubre tanto la parte técnica como la '
     'actualización clínica, esencial cada vez que se publica una nueva versión de las guías.')
make_table(['Concepto', 'Importe/año'],
    [['Actualización de guías y re-indexado', ''],
     ['Revalidación clínica', ''],
     ['Mantenimiento correctivo y evolutivo', ''],
     ['Monitorización de calidad', ''],
     ['Soporte y formación', ''],
     ['Auditoría de seguridad', ''],
     ['Total mantenimiento (sin IVA)', '']],
    widths=[11.5, 5.5])
h2('10.1. Niveles de servicio (SLA)')
para('El contrato incorpora compromisos de nivel de servicio que garantizan la disponibilidad y los '
     'tiempos de atención. Los valores son una propuesta de partida; el cuadro definitivo se fija en el '
     'contrato de servicio.')
make_table(['Indicador', 'Compromiso', 'Medición'],
    [['Disponibilidad', '99,5 % mensual', '24×7, excl. mantenimiento planificado'],
     ['Respuesta — crítica', '2 h laborables', 'Servicio bloqueado'],
     ['Respuesta — mayor', '8 h laborables', 'Funcionalidad degradada'],
     ['Respuesta — menor', '2 días laborables', 'Fallo no bloqueante'],
     ['Resolución — crítica', '1 día laborable', 'Desde el diagnóstico'],
     ['Actualización de guía', '15 días laborables', 'Desde su publicación'],
     ['Soporte', 'L–V 9:00–18:00 (CET)', 'Canal dedicado, ampliable']],
    widths=[5.0, 5.0, 7.0], size=10)

# ========================= 11. CALIDAD =========================
h1('11. Validaciones esperadas')
para('La calidad del asistente se mide con una batería de preguntas clínicas reales antes de su puesta '
     'en producción. La validación final la realiza personal médico; ninguna métrica automática '
     'sustituye esa revisión clínica. Los objetivos de calidad son:')
make_table(['Garantía', 'Objetivo'],
    [['Localiza la sección correcta de la guía', '≥ 90 %'],
     ['Toda afirmación está respaldada y citada', '≥ 95 %'],
     ['La referencia y el nivel de evidencia son correctos', '≥ 95 %'],
     ['No responde cuando la guía no lo cubre', '≈ 100 %'],
     ['Ningún dato sensible se procesa indebidamente', '100 %']],
    widths=[12.5, 4.5])

# ========================= 12. PRÓXIMOS PASOS =========================
h1('12. Próximos pasos')
bullet('Revisión conjunta del alcance, las estimaciones y los supuestos económicos de esta propuesta.')
bullet('Sesión de descubrimiento para cerrar requisitos y criterios clínicos de aceptación.')
bullet('Análisis de protección de datos y regulatorio con la asesoría de la organización (ver Anexo A).')
bullet('Apertura de la cuenta en la nube a nombre de la organización y firma del contrato.')

# ========================= 13. AVISO LEGAL =========================
h1('13. Aviso legal y limitaciones')
for t in [
 'Documento de carácter orientativo; no constituye asesoramiento legal, regulatorio, médico ni financiero.',
 'Los costes de infraestructura y de consumo del modelo de IA son estimaciones de referencia de 2026, se abonan desde una cuenta de la organización y deben verificarse con las tarifas oficiales del proveedor.',
 'El asistente es una herramienta de apoyo y no sustituye el juicio clínico del profesional, responsable último de toda decisión asistencial.',
 'El tratamiento de datos personales de salud está sujeto al RGPD; las bases de licitud y las medidas de protección deben validarse con el Delegado de Protección de Datos y la asesoría jurídica de la organización.',
 'La posible consideración de la herramienta como producto sanitario y las implicaciones del Reglamento de IA de la UE deben evaluarse con asesoría especializada antes de su uso asistencial.',
 'Horas, plazos, métricas y niveles de servicio son previsiones de buena fe y no compromisos contractuales; el alcance definitivo se fija en el contrato.',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t); r.font.size = Pt(9.5); r.font.name=FONT; r.font.color.rgb=BLACK

# ========================= ANEXO A — RGPD =========================
doc.add_page_break()
h1('Anexo A. Análisis de cumplimiento en protección de datos (RGPD)')
para('Documento preparado para su revisión por la asesoría jurídica y el Delegado de Protección de '
     'Datos (DPD) de la organización. No constituye asesoramiento legal; recoge el análisis técnico de '
     'las medidas de protección de datos previstas en la solución, para su validación jurídica.',
     italic=True)

h2('A.1. Objeto y alcance')
para('El presente anexo analiza el tratamiento de datos personales —en particular datos relativos a la '
     'salud— que tiene lugar cuando un médico consulta al asistente aportando información de un paciente '
     'concreto. Su finalidad es facilitar la evaluación jurídica del cumplimiento del Reglamento (UE) '
     '2016/679 (RGPD) y de la Ley Orgánica 3/2018 (LOPDGDD), identificar las medidas técnicas y '
     'organizativas implementadas y señalar los puntos que requieren decisión o validación de la '
     'organización.')

h2('A.2. Marco normativo aplicable')
bullet('Reglamento (UE) 2016/679 (RGPD), en particular los artículos 5 (principios), 6 (licitud), 9 '
       '(categorías especiales de datos), 25 (protección desde el diseño y por defecto), 28 (encargado '
       'del tratamiento), 32 (seguridad), 33–34 (brechas), 35 (evaluación de impacto) y 44–49 '
       '(transferencias internacionales).')
bullet('Ley Orgánica 3/2018 (LOPDGDD) y normativa sectorial sanitaria aplicable.')
bullet('Reglamento (UE) 2024/1689 de Inteligencia Artificial, a efectos de la clasificación de riesgo del sistema.')
bullet('Reglamento (UE) 2017/745 de productos sanitarios (MDR), a efectos de evaluar si la herramienta '
       'puede tener la consideración de producto sanitario.')

h2('A.3. Naturaleza de los datos tratados')
para('Cuando la consulta se formula en abstracto (sin datos de un paciente), no hay tratamiento de datos '
     'personales. Cuando el médico aporta el caso de un paciente, pueden concurrir:')
bullet('identificadores directos (nombre, número de historia clínica, documento '
       'de identidad, fechas concretas, teléfono), que el sistema no necesita y elimina en la entrada.',
       bold_lead='Datos identificativos: ')
bullet('datos relativos a la salud —categoría especial del art. 9 RGPD— (función '
       'renal, comorbilidades, medicación, analíticas), que son los estrictamente necesarios para '
       'afinar la recomendación.', bold_lead='Datos de salud: ')
para('El sistema está diseñado para tratar únicamente los datos clínicos imprescindibles y nunca la '
     'identidad de la persona (principio de minimización, art. 5.1.c).')

h2('A.4. Roles y responsabilidades')
bullet('la organización sanitaria, que decide los fines y medios del tratamiento.', bold_lead='Responsable del tratamiento: ')
bullet('[Tu empresa], que trata los datos por cuenta de la organización para '
       'operar el asistente, mediante un contrato de encargo de tratamiento (art. 28 RGPD).',
       bold_lead='Encargado del tratamiento: ')
bullet('los proveedores de nube e inteligencia artificial (Amazon Web Services '
       'y/o Microsoft Azure), vinculados mediante los correspondientes acuerdos de tratamiento de datos '
       '(DPA) y cláusulas de subencargo.', bold_lead='Subencargados: ')

h2('A.5. Bases de licitud')
para('La determinación de la base de licitud corresponde a la organización con su asesoría jurídica. '
     'Para el tratamiento de datos de salud deberá concurrir, además de una base del art. 6, una '
     'excepción del art. 9.2 RGPD. Las vías habituales en el contexto asistencial son:')
bullet('art. 9.2.h (fines de medicina preventiva o laboral, diagnóstico médico, prestación de '
       'asistencia sanitaria), en relación con el art. 6.1 y la normativa sanitaria.')
bullet('alternativamente, cuando proceda, art. 9.2.a (consentimiento explícito del interesado).')
para('Debe documentarse la base elegida y su encaje con la finalidad de apoyo a la decisión clínica.')

h2('A.6. Minimización y seudonimización')
para('La seudonimización (art. 4.5 y art. 32 RGPD) se aplica en la pasarela de entrada, antes de '
     'cualquier otro tratamiento y antes de que el modelo reciba dato alguno:')
bullet('los identificadores directos se detectan automáticamente y se sustituyen por marcadores '
       'neutros; el modelo y el resto del sistema no llegan a recibirlos.')
bullet('se vigilan las combinaciones de cuasi-identificadores que, aun sin nombre, podrían señalar a una '
       'persona; ante un riesgo elevado de reidentificación, el sistema generaliza esos datos o se '
       'abstiene de procesar la consulta, advirtiendo al médico.')
bullet('el registro de auditoría almacena siempre la versión ya seudonimizada de la consulta, nunca los '
       'identificadores originales.')
para('De este modo, la protección de datos no es un añadido posterior, sino una condición previa al '
     'funcionamiento del asistente (protección desde el diseño y por defecto, art. 25).')

h2('A.7. Residencia de datos y transferencias internacionales')
para('La infraestructura y los puntos de acceso a los modelos se ubican en regiones de la Unión Europea, '
     'de modo que el tratamiento ordinario ocurre dentro del EEE y no implica, por defecto, una '
     'transferencia internacional. En caso de que cualquier componente o soporte pudiera implicar acceso '
     'desde fuera del EEE, deberán aplicarse las garantías del Capítulo V del RGPD (cláusulas '
     'contractuales tipo y, en su caso, evaluación de impacto de las transferencias), lo que se '
     'verificará con la asesoría jurídica antes de la puesta en producción.')

h2('A.8. Garantías de los proveedores de modelos')
para('La elección entre las dos opciones de la sección 9.3 no altera el modelo de protección de datos: '
     'en ambos casos el modelo solo recibe datos ya seudonimizados y se invoca en la UE. Las garantías '
     'relevantes de cada proveedor (a confirmar con la documentación contractual vigente en el momento '
     'de la contratación) son:')
para('A.8.1. Amazon Bedrock', bold=True, size=10.5)
bullet('los datos de entrada y salida no se utilizan para entrenar ni mejorar los modelos de base.')
bullet('los prompts y las respuestas no se comparten con los proveedores de los modelos de base.')
bullet('los datos no se conservan tras el procesamiento de la solicitud.')
bullet('disponibilidad en regiones de la UE; cifrado en tránsito y en reposo; integración con redes '
       'privadas (VPC/PrivateLink) y gestión de claves (KMS). AWS actúa como subencargado bajo su DPA.')
para('A.8.2. Azure OpenAI', bold=True, size=10.5)
bullet('los datos no se utilizan para entrenar ni mejorar los modelos de OpenAI ni de Microsoft.')
bullet('los datos no se comparten con OpenAI; el servicio lo presta Microsoft dentro de su nube.')
bullet('residencia de datos en la UE mediante el despliegue en regiones/zonas de datos europeas.')
bullet('posibilidad de solicitar la exclusión del registro para la monitorización de abuso ("abuse '
       'monitoring opt-out"), de modo que los prompts no se almacenen para revisión humana. Microsoft '
       'actúa como subencargado bajo su DPA.')
para('En ambos casos, la combinación de (i) seudonimización previa, (ii) inferencia en la UE, (iii) '
     'ausencia de entrenamiento con los datos del cliente y (iv) ausencia de retención de prompts es lo '
     'que hace viable, desde la perspectiva del RGPD, el tratamiento de datos específicos del paciente. '
     'Estas condiciones deben quedar reflejadas en los contratos de encargo y subencargo.')

h2('A.9. Medidas de seguridad (art. 32)')
bullet('Cifrado de los datos en tránsito y en reposo (gestión de claves dedicada).')
bullet('Control de acceso por identidad y mínimo privilegio; autenticación de aplicaciones y usuarios a través de la pasarela API.')
bullet('Aislamiento de red (redes privadas y puntos de acceso privados).')
bullet('Registro de auditoría de cada consulta (en su versión seudonimizada).')
bullet('Copias de seguridad y procedimientos de recuperación (ver SLA, sección 10.1).')

h2('A.10. Evaluación de impacto (EIPD/DPIA, art. 35)')
para('Dado que se tratan datos de salud (categoría especial) y potencialmente a gran escala, se '
     'recomienda realizar una Evaluación de Impacto relativa a la Protección de Datos antes de la puesta '
     'en producción, que valore los riesgos de reidentificación, las medidas de mitigación y el riesgo '
     'residual. [Tu empresa] aportará la información técnica necesaria; la EIPD es responsabilidad del '
     'responsable del tratamiento.')

h2('A.11. Derechos de los interesados')
para('Al tratarse de datos seudonimizados y dado que el asistente no constituye el repositorio principal '
     'de la historia clínica, el ejercicio de derechos (acceso, rectificación, supresión, limitación, '
     'oposición) se canaliza a través de los sistemas asistenciales de la organización. El registro de '
     'auditoría se limita a la consulta seudonimizada y a la trazabilidad técnica, y se somete a las '
     'políticas de conservación que defina la organización.')

h2('A.12. Conservación y supresión')
para('Los registros de auditoría se conservan durante el plazo que determine la organización conforme a '
     'la normativa aplicable y a sus políticas internas, transcurrido el cual se suprimen de forma '
     'segura. Los datos de paciente seudonimizados no se conservan más allá de lo necesario para la '
     'trazabilidad de la consulta.')

h2('A.13. Gestión de brechas de seguridad (art. 33–34)')
para('Se establecerán procedimientos de detección, notificación y respuesta ante incidentes de '
     'seguridad, incluida la notificación al responsable del tratamiento sin dilación indebida para que '
     'éste, en su caso, notifique a la autoridad de control y a los interesados en los plazos legales.')

h2('A.14. Producto sanitario y Reglamento de IA')
para('Debe evaluarse, con asesoría especializada, si la herramienta —en tanto sistema de apoyo a la '
     'decisión clínica— puede tener la consideración de producto sanitario conforme al MDR, así como su '
     'clasificación de riesgo bajo el Reglamento de IA de la UE y las obligaciones derivadas. El diseño '
     'del asistente incorpora ya supervisión humana efectiva y trazabilidad, elementos alineados con '
     'dichas exigencias.')

h2('A.15. Puntos a validar por la asesoría jurídica')
for t in [
 'Confirmar la base de licitud (art. 6 y art. 9.2 RGPD) y documentarla.',
 'Firmar el contrato de encargo de tratamiento (art. 28) y verificar las cláusulas de subencargo con AWS/Azure.',
 'Verificar los DPA de los proveedores y las condiciones de no entrenamiento, no retención y residencia en la UE.',
 'Decidir y solicitar, en su caso, la exclusión del registro de monitorización de abuso en Azure OpenAI.',
 'Realizar la EIPD y, si procede, la evaluación de transferencias internacionales.',
 'Definir plazos de conservación y procedimientos de ejercicio de derechos y de gestión de brechas.',
 'Evaluar la consideración como producto sanitario (MDR) y la clasificación bajo el Reglamento de IA.',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t); r.font.size = Pt(10); r.font.name=FONT; r.font.color.rgb=BLACK

import os
out = 'Propuesta_Comercial_Asistente_VIH.docx'
try:
    doc.save(out)
except PermissionError:
    out = 'Propuesta_Comercial_Asistente_VIH_v2.docx'
    try:
        doc.save(out)
    except PermissionError:
        out = 'Propuesta_Asistente_VIH_final.docx'
        doc.save(out)
print('OK ->', out)
