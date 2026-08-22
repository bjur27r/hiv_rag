"""Conversor del banco de preguntas (.xlsx) -> dataset de evaluacion.

Lee `tests/Banco_preguntas_RAG_GeSIDA.xlsx` (hoja con 505 preguntas) y produce:
  - `artifacts/eval_dataset.jsonl`  (siempre)
  - opcionalmente, lo sube a LangSmith como dataset (--upload)

Cada ejemplo conserva la pregunta como `input` y, como `output`/metadatos de
referencia, las columnas que el banco ya trae (nivel, guias, ruta esperada,
anclaje, requiere HITL...). Esto convierte el banco en el dataset canonico contra
el que se miden los evaluadores por nivel (recall de recuperacion, fidelidad de
cita, correctitud numerica, disparo de HITL).

Uso:
    python -m asistente_vih.eval.build_dataset
    python -m asistente_vih.eval.build_dataset --upload
"""
from __future__ import annotations

import argparse
import json
import os
import re

from openpyxl import load_workbook

from ..config import ARTIFACTS_DIR, EVAL_DATASET_PATH, QUESTION_BANK_XLSX

# Mapeo de cabeceras del Excel -> claves normalizadas del dataset.
COLUMNAS = {
    "ID": "id",
    "Nivel": "nivel",
    "Pregunta": "pregunta",
    "Tipo de razonamiento": "tipo_razonamiento",
    "Guías implicadas": "guias_implicadas",
    "Nº guías": "n_guias",
    "Lógica numérica": "logica_numerica",
    "Entidades clave": "entidades_clave",
    "Ruta de recuperación esperada": "ruta_esperada",
    "Respuesta de referencia (a validar)": "respuesta_referencia",
    "Anclaje (sección/tema)": "anclaje",
    "Requiere HITL": "requiere_hitl",
    "Validación clínica": "validacion_clinica",
}

_SI = {"si", "sí", "true", "1", "x"}


def _norm(v: object) -> str:
    return "" if v is None else str(v).strip()


def _nivel_corto(nivel: str) -> int | None:
    m = re.search(r"Nivel\s*(\d)", nivel)
    return int(m.group(1)) if m else None


def _bool(v: str) -> bool:
    return v.strip().lower() in _SI


def _localizar_hoja(wb):
    """Elige la hoja cuyo encabezado contiene 'Pregunta' (la de datos)."""
    for ws in wb.worksheets:
        encabezados = [_norm(c.value) for c in ws[1]]
        if "Pregunta" in encabezados and "ID" in encabezados:
            return ws
    return wb.worksheets[0]


def leer_banco() -> list[dict]:
    wb = load_workbook(QUESTION_BANK_XLSX, read_only=True, data_only=True)
    ws = _localizar_hoja(wb)
    encabezados = [_norm(c.value) for c in ws[1]]
    idx = {h: i for i, h in enumerate(encabezados)}

    ejemplos: list[dict] = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if fila is None or all(v is None for v in fila):
            continue
        reg: dict = {}
        for cab, clave in COLUMNAS.items():
            reg[clave] = _norm(fila[idx[cab]]) if cab in idx else ""
        if not reg.get("pregunta"):
            continue

        nivel = _nivel_corto(reg["nivel"])
        ejemplos.append(
            {
                "inputs": {"pregunta": reg["pregunta"]},
                "outputs": {
                    "respuesta_referencia": reg["respuesta_referencia"],
                    "anclaje": reg["anclaje"],
                    "requiere_hitl": _bool(reg["requiere_hitl"]),
                },
                "metadata": {
                    "id": reg["id"],
                    "nivel": nivel,
                    "nivel_txt": reg["nivel"],
                    "tipo_razonamiento": reg["tipo_razonamiento"],
                    "guias_implicadas": reg["guias_implicadas"],
                    "n_guias": reg["n_guias"],
                    "logica_numerica": reg["logica_numerica"],
                    "entidades_clave": reg["entidades_clave"],
                    "ruta_esperada": reg["ruta_esperada"],
                    "validacion_clinica": reg["validacion_clinica"],
                },
            }
        )
    wb.close()
    return ejemplos


def escribir_jsonl(ejemplos: list[dict]) -> None:
    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    with open(EVAL_DATASET_PATH, "w", encoding="utf-8") as f:
        for e in ejemplos:
            f.write(json.dumps(e, ensure_ascii=False) + "\n")


def resumen(ejemplos: list[dict]) -> None:
    from collections import Counter

    niveles = Counter(e["metadata"]["nivel"] for e in ejemplos)
    hitl = sum(1 for e in ejemplos if e["outputs"]["requiere_hitl"])
    multi = sum(1 for e in ejemplos if (e["metadata"]["n_guias"] or "0").isdigit()
                and int(e["metadata"]["n_guias"]) >= 2)
    print("\nResumen del dataset de evaluacion")
    print("-" * 40)
    print(f"  Total preguntas: {len(ejemplos)}")
    for niv in sorted(k for k in niveles if k is not None):
        print(f"  Nivel {niv}: {niveles[niv]}")
    print(f"  Multi-guia (>=2): {multi}")
    print(f"  Requiere HITL: {hitl}")
    print("-" * 40)
    print(f"  -> {EVAL_DATASET_PATH}")


def subir_langsmith(ejemplos: list[dict]) -> None:
    try:
        from langsmith import Client
    except ImportError:
        print("  [!] langsmith no instalado; omito subida (pip install langsmith).")
        return
    if not os.getenv("LANGSMITH_API_KEY"):
        print("  [!] Falta LANGSMITH_API_KEY; omito subida.")
        return

    nombre = os.getenv("LANGSMITH_DATASET", "gesida-banco-505")
    client = Client()
    if client.has_dataset(dataset_name=nombre):
        ds = client.read_dataset(dataset_name=nombre)
    else:
        ds = client.create_dataset(
            dataset_name=nombre,
            description="Banco GeSIDA: 505 preguntas, 4 niveles, multi-salto y HITL.",
        )
    client.create_examples(
        inputs=[e["inputs"] for e in ejemplos],
        outputs=[e["outputs"] for e in ejemplos],
        metadata=[e["metadata"] for e in ejemplos],
        dataset_id=ds.id,
    )
    print(f"  [ok] Subidos {len(ejemplos)} ejemplos al dataset '{nombre}'.")


def main() -> None:
    ap = argparse.ArgumentParser(description="Banco de preguntas .xlsx -> dataset de evaluacion")
    ap.add_argument("--upload", action="store_true", help="subir el dataset a LangSmith")
    args = ap.parse_args()

    ejemplos = leer_banco()
    escribir_jsonl(ejemplos)
    resumen(ejemplos)
    if args.upload:
        subir_langsmith(ejemplos)


if __name__ == "__main__":
    main()
