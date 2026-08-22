"""Paso 1 - Export de validacion clinica.

Convierte el banco de 505 preguntas en un libro Excel COMODO para que el equipo
clinico valide las respuestas de referencia (la columna del banco esta "Pendiente").
Es el prerrequisito para poder usar el banco como gating de calidad: sin un patron
oro validado, las metricas son solo aproximadas.

Genera artifacts/validacion_clinica.xlsx con las columnas del banco + columnas
vacias para la revision (Correcta / Correccion / Validador / Fecha / Notas).

    python -m asistente_vih.eval.export_validacion
"""
from __future__ import annotations

import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import ARTIFACTS_DIR, EVAL_DATASET_PATH

SALIDA = ARTIFACTS_DIR / "validacion_clinica.xlsx"

CABECERAS = ["ID", "Nivel", "Pregunta", "Respuesta de referencia (banco)", "Anclaje",
             "Guias implicadas", "Requiere HITL",
             "¿Correcta? (Si/No/Matiz)", "Correccion / matiz", "Validador", "Fecha", "Notas"]


def main() -> None:
    ejemplos = [json.loads(l) for l in open(EVAL_DATASET_PATH, encoding="utf-8") if l.strip()]
    wb = Workbook()
    ws = wb.active
    ws.title = "Validacion"

    # Cabecera con estilo.
    azul = PatternFill("solid", fgColor="1F3B57")
    for c, txt in enumerate(CABECERAS, start=1):
        cel = ws.cell(row=1, column=c, value=txt)
        cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = azul
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"

    # Validacion desplegable Si/No/Matiz en la columna 8.
    dv = DataValidation(type="list", formula1='"Si,No,Matiz"', allow_blank=True)
    ws.add_data_validation(dv)

    for i, e in enumerate(ejemplos, start=2):
        m, o = e["metadata"], e["outputs"]
        fila = [m.get("id"), m.get("nivel"), e["inputs"]["pregunta"],
                o.get("respuesta_referencia"), o.get("anclaje"), m.get("guias_implicadas"),
                "Si" if o.get("requiere_hitl") else "No", "", "", "", "", ""]
        for c, v in enumerate(fila, start=1):
            cel = ws.cell(row=i, column=c, value=v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(ws.cell(row=i, column=8))

    anchos = [8, 8, 50, 55, 16, 24, 8, 14, 40, 14, 12, 30]
    for c, w in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w

    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(SALIDA)
    print(f"OK: {len(ejemplos)} preguntas -> {SALIDA}")
    print("Reparte por nivel y guia entre los revisores; el campo '¿Correcta?' alimenta el gating.")


if __name__ == "__main__":
    main()
