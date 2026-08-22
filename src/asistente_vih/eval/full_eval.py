# -*- coding: utf-8 -*-
"""Evaluacion COMPLETA sobre las 505: genera, verifica y exporta para revision clinica.

Un solo run produce dos entregables:
  - artifacts/full_eval.jsonl  : respuesta real del sistema + veredicto por pregunta.
  - artifacts/revision_respuestas.xlsx : libro para que el equipo clinico puntue.
Y las metricas automaticas por nivel (fidelidad/abstencion).

RESUMIBLE: cada respuesta se guarda al generarse; si se corta, re-ejecutar continua
donde se quedo (salta las ya hechas). Construir el xlsx desde el jsonl parcial:
  python -m asistente_vih.eval.full_eval --solo-xlsx

    python -m asistente_vih.eval.full_eval            # las 505 (largo, caro)
    python -m asistente_vih.eval.full_eval --limite 50
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict

from ..config import ARTIFACTS_DIR, EVAL_DATASET_PATH

JSONL = ARTIFACTS_DIR / "full_eval.jsonl"
XLSX = ARTIFACTS_DIR / "revision_respuestas.xlsx"

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass


def _ids_hechos() -> set[str]:
    """IDs con respuesta CORRECTA (sin error). Las fallidas se reintentan al re-ejecutar."""
    if not JSONL.exists():
        return set()
    ok = set()
    for l in open(JSONL, encoding="utf-8"):
        if not l.strip():
            continue
        r = json.loads(l)
        if not r.get("error") and r.get("respuesta"):
            ok.add(r["id"])
    return ok


def _dedup() -> list[dict]:
    """Un registro por ID, prefiriendo el exitoso (para metricas/xlsx sin duplicar)."""
    best: dict[str, dict] = {}
    for l in open(JSONL, encoding="utf-8"):
        if not l.strip():
            continue
        r = json.loads(l)
        prev = best.get(r["id"])
        if prev is None or (prev.get("error") and not r.get("error")):
            best[r["id"]] = r
    return list(best.values())


def correr(limite: int | None) -> None:
    from ..llm.client import LLM, MODELO_SINTESIS, MODELO_VERIFICADOR
    from ..llm.generate import generar_respuesta
    from ..llm.verify import verificar_fidelidad
    from ..retrieval.dense import DenseRetriever
    from ..retrieval.decompose import DecomposingRetriever

    banco = [json.loads(l) for l in open(EVAL_DATASET_PATH, encoding="utf-8") if l.strip()]
    if limite:
        banco = banco[:limite]
    hechos = _ids_hechos()
    pend = [e for e in banco if e["metadata"].get("id") not in hechos]
    print(f"Total {len(banco)} | ya hechas {len(hechos)} | pendientes {len(pend)}")

    denso = DenseRetriever()
    decomp = DecomposingRetriever(denso)
    sint, verif = LLM(MODELO_SINTESIS), LLM(MODELO_VERIFICADOR)
    print(f"Pipeline: denso + {sint.provider}:{sint.modelo} + {verif.modelo} | "
          f"{'SIMULADO' if sint.simulado else 'REAL'}\n")

    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    with open(JSONL, "a", encoding="utf-8") as out:
        for i, e in enumerate(pend, 1):
            m, o = e["metadata"], e["outputs"]
            q = e["inputs"]["pregunta"]
            multi = (m.get("n_guias") or "1").isdigit() and int(m.get("n_guias")) >= 2
            rec = {"id": m.get("id"), "nivel": m.get("nivel"), "pregunta": q,
                   "guias_implicadas": m.get("guias_implicadas"),
                   "respuesta_referencia": o.get("respuesta_referencia"), "anclaje": o.get("anclaje")}
            try:
                hits = (decomp if multi else denso).search(q, k=20 if multi else 8)
                gen = generar_respuesta(q, hits, llm=sint)
                v = verificar_fidelidad(q, gen["respuesta"], hits, llm=verif)
                rec.update(guias_recuperadas=sorted({h.guia for h in hits}),
                           respuesta=gen["respuesta"], abstenida=gen["abstenida"],
                           veredicto=v.get("veredicto"), fiel=bool(v.get("fiel")),
                           sin_respaldo=v.get("afirmaciones_sin_respaldo", []))
            except Exception as ex:
                rec.update(error=f"{type(ex).__name__}: {ex}", respuesta="", fiel=None)
            out.write(json.dumps(rec, ensure_ascii=False) + "\n")
            out.flush()
            print(f"  [{i}/{len(pend)}] N{rec['nivel']} {rec['id']}: "
                  f"{rec.get('veredicto') or rec.get('error','?')}", flush=True)


def metricas() -> None:
    rows = _dedup()
    agg = defaultdict(lambda: {"n": 0, "fiel": 0, "abst": 0, "parcial": 0, "err": 0})
    for r in rows:
        a = agg[r.get("nivel")]
        a["n"] += 1
        if r.get("error"):
            a["err"] += 1; continue
        a["fiel"] += int(bool(r.get("fiel")))
        a["abst"] += int(bool(r.get("abstenida")))
        a["parcial"] += int(r.get("veredicto") == "parcial")
    print("\n=== Metricas a nivel de respuesta (n total = {}) ===".format(len(rows)))
    tot = defaultdict(int)
    for niv in sorted(x for x in agg if x is not None):
        a = agg[niv]
        for k in a: tot[k] += a[k]
        base = max(a["n"] - a["err"], 1)
        print(f"  Nivel {niv} (n={a['n']}, err={a['err']}): fidelidad={a['fiel']/base:.0%} "
              f"abstencion={a['abst']/base:.0%} parcial={a['parcial']/base:.0%}")
    base = max(tot["n"] - tot["err"], 1)
    print(f"  GLOBAL (n={tot['n']}): fidelidad={tot['fiel']/base:.0%} abstencion={tot['abst']/base:.0%}")


def construir_xlsx() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = sorted(_dedup(), key=lambda r: r.get("id") or "")
    wb = Workbook(); ws = wb.active; ws.title = "Revision"
    cab = ["ID", "Nivel", "Pregunta", "Respuesta del SISTEMA", "Veredicto fidelidad",
           "Respuesta referencia (banco)", "Anclaje", "Guias recuperadas",
           "¿Correcta clinicamente? (Si/No/Matiz)", "Comentario / correccion", "Validador"]
    azul = PatternFill("solid", fgColor="1F3B57")
    for c, t in enumerate(cab, 1):
        cel = ws.cell(1, c, t); cel.font = Font(bold=True, color="FFFFFF", size=10)
        cel.fill = azul; cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"Si,No,Matiz"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, r in enumerate(rows, 2):
        vals = [r.get("id"), r.get("nivel"), r.get("pregunta"),
                r.get("respuesta") or f"[ERROR] {r.get('error','')}", r.get("veredicto"),
                r.get("respuesta_referencia"), r.get("anclaje"),
                ", ".join(r.get("guias_recuperadas", [])), "", "", ""]
        for c, v in enumerate(vals, 1):
            ws.cell(i, c, v).alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(ws.cell(i, 9))
    for c, w in enumerate([8, 7, 45, 70, 12, 45, 14, 22, 16, 40, 12], 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    wb.save(XLSX)
    print(f"\nOK xlsx: {len(rows)} filas -> {XLSX}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--limite", type=int)
    ap.add_argument("--solo-xlsx", action="store_true", help="solo (re)construye el xlsx desde el jsonl")
    args = ap.parse_args()
    if not args.solo_xlsx:
        correr(args.limite)
    metricas()
    construir_xlsx()


if __name__ == "__main__":
    main()
